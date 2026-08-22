from __future__ import annotations

import asyncio
import os
import re
import tempfile
import threading
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UPLOAD_FILES = 5
MAX_AGGREGATE_BYTES = 25 * 1024 * 1024
EXTRACTION_TIMEOUT_SECONDS = 60
ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".mp3"}
ALLOWED_MIME_TYPES = {
    ".txt": {"text/plain", "application/octet-stream"},
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/zip",
        "application/octet-stream",
    },
    ".mp3": {"audio/mpeg", "audio/mp3", "application/octet-stream"},
}


class ExtractionError(ValueError):
    pass


class SpeechUnavailableError(RuntimeError):
    pass


def sanitize_filename(filename: str) -> str:
    safe = Path(filename).name
    safe = re.sub(r"[\x00-\x1f\x7f]", "", safe).strip()
    return safe[:160] or "upload"


def speech_is_configured() -> bool:
    return bool(
        os.environ.get("AZURE_SPEECH_KEY")
        and os.environ.get("AZURE_SPEECH_REGION")
    )


def validate_file(filename: str, content_type: str | None, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ExtractionError(
            "지원하지 않는 파일입니다. TXT, MD, PDF, DOCX, MP3만 사용할 수 있습니다."
        )
    normalized_type = (content_type or "application/octet-stream").split(";")[0]
    if normalized_type not in ALLOWED_MIME_TYPES[extension]:
        raise ExtractionError(
            f"{extension.upper()} 확장자와 MIME 유형이 일치하지 않습니다."
        )
    if not data:
        raise ExtractionError("파일이 비어 있습니다.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExtractionError("파일은 최대 10MB까지 업로드할 수 있습니다.")
    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise ExtractionError("유효한 PDF 서명이 없습니다.")
    if extension == ".docx":
        if not data.startswith(b"PK"):
            raise ExtractionError("유효한 DOCX ZIP 서명이 없습니다.")
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ExtractionError("DOCX 문서 본문을 찾을 수 없습니다.")
                if any(name.endswith("vbaProject.bin") for name in archive.namelist()):
                    raise ExtractionError("매크로가 포함된 문서는 지원하지 않습니다.")
        except zipfile.BadZipFile as exc:
            raise ExtractionError("손상된 DOCX 파일입니다.") from exc
    if extension == ".mp3" and not (
        data.startswith(b"ID3")
        or (len(data) > 1 and data[0] == 0xFF and data[1] & 0xE0 == 0xE0)
    ):
        raise ExtractionError("유효한 MP3 오디오 서명이 없습니다.")
    return extension


def extract_text(data: bytes) -> str:
    if b"\x00" in data:
        raise ExtractionError("텍스트 파일에 바이너리 데이터가 포함되어 있습니다.")
    for encoding in ("utf-8-sig", "utf-8", "cp949"):
        try:
            text = data.decode(encoding)
            if text.strip():
                return text.strip()
        except UnicodeDecodeError:
            continue
    raise ExtractionError("UTF-8 또는 CP949 텍스트로 해석할 수 없습니다.")


def extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception as exc:
        raise ExtractionError("PDF 텍스트를 읽지 못했습니다.") from exc
    if len(text) < 20:
        raise ExtractionError(
            "텍스트가 없는 이미지형 PDF입니다. OCR은 지원하지 않으므로 "
            "텍스트 PDF 또는 TXT로 내보내 주세요."
        )
    return text


def extract_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise ExtractionError("DOCX 내용을 읽지 못했습니다.") from exc
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(block for block in blocks if block).strip()
    if len(text) < 20:
        raise ExtractionError("DOCX에서 분석할 텍스트를 찾지 못했습니다.")
    return text


def _transcribe_mp3_sync(data: bytes) -> str:
    if not speech_is_configured():
        raise SpeechUnavailableError(
            "MP3 전사는 현재 Azure Speech 설정이 없어 사용할 수 없습니다. "
            "TXT/PDF/DOCX 샘플 경로는 계속 사용할 수 있습니다."
        )
    import azure.cognitiveservices.speech as speechsdk

    speech_config = speechsdk.SpeechConfig(
        subscription=os.environ["AZURE_SPEECH_KEY"],
        region=os.environ["AZURE_SPEECH_REGION"],
    )
    speech_config.speech_recognition_language = os.environ.get(
        "AZURE_SPEECH_LANGUAGE", "ko-KR"
    )
    chunks: list[str] = []
    done = threading.Event()
    temporary_path = ""
    try:
        with tempfile.NamedTemporaryFile(suffix=".mp3", delete=False) as temporary:
            temporary.write(data)
            temporary_path = temporary.name
        audio_config = speechsdk.audio.AudioConfig(filename=temporary_path)
        recognizer = speechsdk.SpeechRecognizer(
            speech_config=speech_config, audio_config=audio_config
        )
        recognizer.recognized.connect(
            lambda event: chunks.append(event.result.text)
            if event.result.text
            else None
        )
        recognizer.session_stopped.connect(lambda _: done.set())
        recognizer.canceled.connect(lambda _: done.set())
        recognizer.start_continuous_recognition_async().get()
        if not done.wait(EXTRACTION_TIMEOUT_SECONDS):
            recognizer.stop_continuous_recognition_async().get()
            raise ExtractionError("MP3 전사가 60초 제한을 초과했습니다.")
        recognizer.stop_continuous_recognition_async().get()
    finally:
        if temporary_path:
            Path(temporary_path).unlink(missing_ok=True)
    text = " ".join(chunks).strip()
    if not text:
        raise ExtractionError("MP3에서 음성을 인식하지 못했습니다.")
    return text


async def extract_upload(
    filename: str, content_type: str | None, data: bytes
) -> tuple[str, bool, str]:
    extension = validate_file(filename, content_type, data)
    if extension in {".txt", ".md"}:
        return extract_text(data), False, "텍스트를 안전하게 읽었습니다."
    if extension == ".pdf":
        return await asyncio.to_thread(extract_pdf, data), False, "PDF 텍스트를 추출했습니다."
    if extension == ".docx":
        return (
            await asyncio.to_thread(extract_docx, data),
            False,
            "DOCX 문단과 표 텍스트를 추출했습니다.",
        )
    transcription = await asyncio.to_thread(_transcribe_mp3_sync, data)
    editable = (
        f"발표자: {transcription}\n"
        "청중: 질문 또는 피드백을 입력해 주세요.\n"
        "발표자: 답변을 입력해 주세요."
    )
    return (
        editable,
        True,
        "화자 분리는 적용하지 않았습니다. 분석 전에 화자 라벨을 확인해 주세요.",
    )
