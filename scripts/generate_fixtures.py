from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from app.samples import load_samples

OUTPUT_DIR = Path(__file__).parent.parent / "fixtures"


def wrap(text: str, width: int = 44) -> list[str]:
    lines: list[str] = []
    for source_line in text.splitlines():
        current = ""
        for character in source_line:
            current += character
            if len(current) >= width:
                lines.append(current)
                current = ""
        if current:
            lines.append(current)
    return lines


def main() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)
    samples = {sample.id: sample for sample in load_samples()}
    meeting = samples["sprint-planning"]
    presentation = samples["product-pitch"]

    (OUTPUT_DIR / "meeting-sprint.txt").write_text(
        meeting.transcript, encoding="utf-8"
    )

    document = Document()
    document.add_heading(presentation.title, level=1)
    document.add_paragraph(presentation.description)
    for line in presentation.transcript.splitlines():
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "발표자"
    table.cell(0, 1).text = presentation.recommended_self
    table.cell(1, 0).text = "분석 모드"
    table.cell(1, 1).text = "발표 개선"
    document.save(OUTPUT_DIR / "presentation-pitch.docx")

    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    pdf = canvas.Canvas(
        str(OUTPUT_DIR / "presentation-pitch.pdf"),
        pageCompression=1,
    )
    pdf.setTitle(presentation.title)
    pdf.setFont("HYSMyeongJo-Medium", 10)
    y = 800
    for line in wrap(presentation.transcript):
        if y < 50:
            pdf.showPage()
            pdf.setFont("HYSMyeongJo-Medium", 10)
            y = 800
        pdf.drawString(42, y, line)
        y -= 16
    pdf.save()


if __name__ == "__main__":
    main()
